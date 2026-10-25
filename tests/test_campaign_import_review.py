import csv
import io

import pytest
from openpyxl import Workbook,load_workbook
from test_workflow import app,token,register
from test_send import setup_send
from test_message_composer import query


def csv_upload():
    data=io.StringIO();writer=csv.writer(data)
    writer.writerow(['name','phone','info1'])
    writer.writerow(['=1+1','bad,0901234567; +84 912 345 678','Gold'])
    writer.writerow(['Binh','0912345678','Silver'])
    writer.writerow(['Invalid','bad','Bronze'])
    writer.writerow(['Loan','0934567890','Platinum'])
    return io.BytesIO(data.getvalue().encode()),'customers.csv'


def request(client,path='/campaigns/1/data/preview',**extra):
    return client.post(path,data=dict(csrf=token(client),mode='file',file=csv_upload(),**extra))


@pytest.mark.parametrize('policy,valid,duplicate,excluded',[
    ('first',3,0,1),('all',3,1,1),('skip',2,0,2)])
def test_multi_phone_policies_keep_metadata_and_report_selected_excluded_duplicates(app,policy,valid,duplicate,excluded):
    client,connector=setup_send(app)
    data=request(client,multi_phone=policy).json
    assert len(data['rows'])==valid and data['duplicates']==duplicate and len(data['excluded'])==excluded
    assert len(data['duplicate_rows'])==duplicate
    assert data['source_rows']==4 and data['items_count']==valid+duplicate+excluded
    if policy=='all':
        shared=[row for row in data['rows'] if row['phone'] in ('0901234567','0912345678')]
        assert all(row['name']=='=1+1' and row['info1']=='Gold' for row in shared)
        assert data['duplicate_rows'][0]['name']=='Binh'
    if policy=='first':
        assert data['ignored_phones']==1 and data['ignored_invalid']==1
    if policy=='skip':
        assert data['skipped_rows']==1 and 'nhiều SĐT' in data['excluded'][0]['reason']
    assert query(app,'SELECT count(*) FROM contacts')==[(1,)]
    assert query(app,'SELECT count(*) FROM campaign_contacts')==[(0,)]
    assert not connector.message_action.called


def excel_upload():
    book=Workbook();first=book.active;first.title='Main';first.append(['phone']);first.append(['0901234567'])
    sheet=book.create_sheet('CSKH');sheet.append(['MST','Tên','Số điện thoại']);sheet.append(['A','Lan','0912345678'])
    output=io.BytesIO();book.save(output);output.seek(0);return output,'multi.xlsx'


def test_inspection_reveals_sheets_columns_and_source_without_validation_or_writes(app):
    client,_=setup_send(app)
    result=client.post('/campaigns/1/data/preview',data=dict(csrf=token(client),mode='file',phase='inspect',file=excel_upload())).json
    assert result['sheets']==['Main','CSKH'] and result['sheet']=='Main'
    assert 'rows' not in result and result['columns']==[dict(value='0',label='A')]
    result=client.post('/campaigns/1/data/preview',data=dict(csrf=token(client),mode='file',phase='inspect',sheet='CSKH',file=excel_upload())).json
    assert result['sheet']=='CSKH' and result['mapping']['phone']=='2' and result['sample'][1][1]=='Lan'
    assert [col['label'] for col in result['columns']]==['A','B','C']
    assert client.post('/campaigns/1/data/preview',data=dict(csrf=token(client),mode='file',sheet='Unknown',file=excel_upload())).status_code==400
    assert query(app,'SELECT count(*) FROM contacts')==[(1,)]


@pytest.mark.parametrize('category,count',[('valid',3),('invalid',1),('duplicates',1)])
def test_excel_exports_match_result_groups_preserve_zero_and_never_make_formulas(app,category,count):
    client,_=setup_send(app)
    response=request(client,'/campaigns/1/data/export',multi_phone='all',category=category)
    assert response.status_code==200 and '.xlsx' in response.headers['Content-Disposition']
    book=load_workbook(io.BytesIO(response.data),data_only=False);sheet=book.active
    assert sheet.max_row==count+1
    assert all(cell.data_type!='f' for row in sheet for cell in row)
    if category=='valid':
        assert sheet['B2'].value=='=1+1' and sheet['B2'].data_type=='s'
        assert sheet['C2'].value=='0901234567' and sheet['C2'].data_type=='s'
    assert query(app,'SELECT count(*) FROM campaign_events WHERE kind="data"')==[(0,)]


def test_export_is_owner_scoped_csrf_protected_and_policies_are_validated(app):
    client,_=setup_send(app);other=app.test_client();register(other,'other')
    assert request(other,'/campaigns/1/data/export',category='valid').status_code==404
    assert client.post('/campaigns/1/data/export',data=dict(category='valid')).status_code==400
    assert request(client,multi_phone='unknown').status_code==400
    assert request(client,'/campaigns/1/data/export',category='unknown').status_code==400
