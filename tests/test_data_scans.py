import io
import json
import sqlite3
from unittest.mock import Mock

from openpyxl import load_workbook
from test_workflow import app,post,register,token


def query(app,sql,args=()):
    with sqlite3.connect(app.config['DATABASE']) as db:return db.execute(sql,args).fetchall()


def scan_page(kind='friends',offset=0,**changes):
    data=dict(state='ok',rows=[dict(uid='123',name='Lan',phone='0901234567',avatar='https://avatar.zadn.vn/lan.jpg',role='friend',kind='user',member_count=0),
        dict(uid='456',name='UID only',phone='09*****678',avatar='http://127.0.0.1/private',role='member',kind='user',member_count=0)],
        has_more=False,next_offset=offset+100,total=2,limited=False,source_name='Bạn bè')
    if kind=='groups':data.update(rows=[dict(uid='900',name='Group A',phone='',avatar='',role='group',kind='group',member_count=10),dict(uid='901',name='Group B',phone='',avatar='',role='group',kind='group',member_count=20)],source_name='Group đã tham gia')
    data.update(changes);return data


def setup_scan(app):
    client=app.test_client();register(client);post(client,'/accounts',{'name':'Viettax'});post(client,'/data-groups',{'name':'Khách hàng'})
    app.config['ZALO_QR_ENABLED']=True
    connector=app.extensions['zalo_connector'];connector.status=Mock(return_value={'state':'connected'})
    connector.scan_action=Mock(side_effect=lambda owner,account,kind,offset=0,group_id='':scan_page(kind,offset))
    connector.message_action=Mock(return_value={'state':'resolved','uid':'123'})
    return client,connector


def start(client,kind='friends',**fields):
    return post(client,'/data-scans/start',{'kind':kind,'account_id':'1',**fields})


def test_scan_retains_uid_only_scrubs_fields_and_does_not_send_or_import(app):
    client,connector=setup_scan(app)
    data=scan_page();data['rows'][0].update(cookie='PRIVATE_COOKIE',token='PRIVATE_TOKEN');data['cookie']='PRIVATE_COOKIE'
    connector.scan_action.return_value=data;connector.scan_action.side_effect=None
    response=start(client);assert response.status_code==200
    result=client.get('/data-scans/1/rows').json
    assert result['loaded']==2 and result['rows'][0]['phone']=='0901234567'
    assert result['rows'][1]['phone']=='' and result['rows'][1]['uid']=='456'
    assert result['rows'][1]['avatar']==''
    assert 'PRIVATE_' not in json.dumps(result)
    assert query(app,'SELECT count(*) FROM contacts')==[(0,)]
    assert not connector.message_action.called
    assert client.get('/data-scans/1/rows?phone=yes').json['total']==1
    assert len(client.get('/data-scans/1/ids').json['ids'])==2
    assert len(client.get('/data-scans/1/phone-ids').json['ids'])==1


def test_scan_owner_csrf_and_account_boundaries(app):
    client,connector=setup_scan(app);start(client)
    other=app.test_client();register(other,'other')
    for path in ('/data-scans/1/rows','/data-scans/1/ids','/data-scans/1/phone-ids'):
        assert other.get(path).status_code==404
    for path in ('/data-scans/1/export','/data-scans/1/import','/data-scans/1/more','/data-scans/1/check-phone'):
        assert post(other,path,{'row_id':'1','group_id':'1'}).status_code==404
        assert client.post(path).status_code==400
    assert post(other,'/data-scans/start',{'account_id':'1','kind':'friends'}).status_code==404
    assert post(other,'/data-scans/groups',{'account_id':'1'}).status_code==404
    assert len(client.get('/data-scans/catalog').json['scans'])==1
    assert other.get('/data-scans/catalog').json['scans']==[]


def test_scan_phone_import_keeps_uid_only_in_snapshot_and_existing_profile_intact(app):
    client,_=setup_scan(app)
    post(client,'/contacts',{'name':'Original','phone':'0901234567','tag':'Original group','consent':'on'})
    start(client)
    ids=client.get('/data-scans/1/ids').json['ids']
    result=post(client,'/data-scans/1/import',{'group_id':'1','ids':json.dumps(ids)}).json
    assert result['eligible']==1 and result['skipped']==1 and result['added']==1
    assert query(app,'SELECT name,phone,tag,consent FROM contacts')==[('Original','0901234567','Original group',1)]
    assert query(app,"SELECT label FROM contact_data_sources WHERE source_type='scan'")==[('Quét Zalo · Bạn bè',)]
    assert client.get('/data-scans/1/rows').json['loaded']==2
    again=post(client,'/data-scans/1/import',{'group_id':'1','ids':json.dumps(ids)}).json
    assert again['added']==0
    assert query(app,'SELECT count(*) FROM contact_group_memberships WHERE group_name=?',('Khách hàng',))==[(1,)]


def test_scan_export_keeps_phone_zero_and_formula_like_names_as_text(app):
    client,connector=setup_scan(app);data=scan_page();data['rows'][0]['name']='=1+1';connector.scan_action.side_effect=None;connector.scan_action.return_value=data
    start(client)
    response=post(client,'/data-scans/1/export')
    assert response.status_code==200
    sheet=load_workbook(io.BytesIO(response.data)).active
    assert sheet.max_row==3 and sheet['B2'].value=='=1+1' and sheet['B2'].data_type=='s'
    assert sheet['C2'].value=='0901234567' and sheet['D3'].value=='456'


def test_scan_partial_error_retry_preserves_previous_rows_and_never_exposes_error(app):
    client,connector=setup_scan(app)
    connector.scan_action.side_effect=[scan_page(has_more=True),ValueError('PRIVATE_COOKIE'),scan_page(offset=100,rows=[dict(uid='789',name='Next',phone='',kind='user')])]
    start(client)
    result=post(client,'/data-scans/1/more');assert result.status_code==400 and 'PRIVATE_' not in result.text
    assert client.get('/data-scans/1/rows').json['loaded']==2
    assert post(client,'/data-scans/1/more').status_code==200
    result=client.get('/data-scans/1/rows').json
    assert result['loaded']==3 and not result['scan']['has_more']
    assert query(app,'SELECT count(*) FROM contacts')==[(0,)]


def test_phone_check_is_read_only_and_uid_mismatch_cannot_be_imported(app):
    client,connector=setup_scan(app);start(client)
    connector.message_action.return_value={'state':'resolved','uid':'999'}
    assert post(client,'/data-scans/1/check-phone',{'row_id':'1'}).json['state']=='mismatch'
    assert connector.message_action.call_args.args[-1]=='resolve'
    assert post(client,'/data-scans/1/import',{'group_id':'1'}).status_code==400
    assert post(client,'/data-scans/1/check-phone',{'row_id':'2'}).status_code==400
    assert query(app,'SELECT count(*) FROM contacts')==[(0,)]


def test_group_scan_cannot_import_groups_as_phone_contacts_and_rejects_foreign_row_ids(app):
    client,_=setup_scan(app);start(client,'groups');start(client)
    assert post(client,'/data-scans/1/import',{'group_id':'1'}).status_code==400
    assert post(client,'/data-scans/2/import',{'group_id':'1','ids':'[1]'}).status_code==404
    assert query(app,'SELECT count(*) FROM contacts')==[(0,)]


def test_history_uses_successful_local_sends_and_does_not_call_zalo(app):
    from test_send import setup_send,prepare
    client,connector=setup_send(app);prepare(client)
    connector.message_action.return_value={'state':'sent','message_id':'1'}
    post(client,'/campaigns/1/send/1',{'confirm_send':'yes'})
    connector.scan_action=Mock();connector.message_action.reset_mock();connector.status.return_value={'state':'disconnected'}
    result=start(client,'history')
    assert result.status_code==200
    rows=client.get('/data-scans/1/rows').json['rows']
    assert len(rows)==1 and rows[0]['phone']=='0901234567'
    assert not connector.scan_action.called and not connector.message_action.called


def test_read_only_scan_can_resume_after_interrupted_worker_lease_expires(app):
    client,connector=setup_scan(app)
    connector.scan_action.side_effect=[scan_page(has_more=True),scan_page(offset=100,rows=[dict(uid='789',name='Next',phone='',kind='user')])]
    start(client)
    with sqlite3.connect(app.config['DATABASE']) as db:
        db.execute("UPDATE data_scans SET status='loading',loading_at=CURRENT_TIMESTAMP WHERE id=1")
    assert post(client,'/data-scans/1/more').status_code==409
    with sqlite3.connect(app.config['DATABASE']) as db:
        db.execute("UPDATE data_scans SET loading_at=datetime('now','-3 minutes') WHERE id=1")
    assert post(client,'/data-scans/1/more').status_code==200
    assert client.get('/data-scans/1/rows').json['loaded']==3
    assert query(app,'SELECT count(*) FROM contacts')==[(0,)]
