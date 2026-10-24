import io
import sqlite3

from openpyxl import Workbook
from test_workflow import app, post, register


def database(app):
    return sqlite3.connect(app.config['DATABASE'])


def test_group_import_metadata_duplicates_and_consent(app):
    client = app.test_client()
    register(client)
    assert 'Nhóm data' in post(client, '/data-groups', {'name': 'VIP'}).text
    result = post(client, '/data-groups/import', {'group': 'VIP', 'consent': 'yes',
        'file': (io.BytesIO('phone,name,email,info1\n+84901234567,Lan,lan@example.com,A\n'.encode()), 'contacts.csv')})
    assert 'Lan' in result.text
    post(client, '/data-groups', {'name': 'Other'})
    post(client, '/data-groups/import', {'group': 'Other', 'phones': '0901234567\n0912345678', 'consent': ''})
    with database(app) as db:
        assert db.execute('SELECT name,tag,consent,email,info1 FROM contacts WHERE phone=?', ('0901234567',)).fetchone() == ('Lan', 'VIP', 1, 'lan@example.com', 'A')
        assert db.execute('SELECT consent FROM contacts WHERE phone=?', ('0912345678',)).fetchone() == (0,)
    # Existing profiles are shared with the destination; their consent is retained.
    with database(app) as db:
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='Other'").fetchone() == (2,)
    post(client, '/campaigns', {'name': 'Group', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Other'})
    with database(app) as db:
        assert db.execute('SELECT audience_mode,tag FROM campaigns').fetchone() == ('group', 'Other')


def test_campaign_file_import_edit_and_snapshot_group(app):
    client = app.test_client()
    register(client)
    book = Workbook()
    book.active.append(['phone', 'name', 'info2'])
    book.active.append([901234567, 'Lan', 'Excel'])
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    result = post(client, '/campaigns', {'name': 'Excel', 'message': 'Hi', 'data_mode': 'direct',
        'data_file': (stream, 'data.xlsx'), 'phones_consent': 'yes'})
    assert 'Excel' in result.text
    with database(app) as db:
        assert db.execute('SELECT name,phone,info2 FROM contacts').fetchone() == ('Lan', '0901234567', 'Excel')
        assert db.execute('SELECT audience_mode FROM campaigns').fetchone() == ('list',)
    assert 'Cách thêm data vào chiến dịch' in client.get('/campaigns/1/edit').text
    post(client, '/data-groups', {'name': 'VIP'})
    post(client, '/data-groups/import', {'group': 'VIP', 'phones': '0912345678', 'consent': 'yes'})
    post(client, '/campaigns/1/edit', {'name': 'Snapshot', 'message': 'Hi', 'data_mode': 'direct',
        'source_tags': 'VIP', 'phones_consent': 'yes'})
    post(client, '/data-groups/import', {'group': 'VIP', 'phones': '0923456789', 'consent': 'yes'})
    with database(app) as db:
        assert db.execute('SELECT p.phone FROM contacts p JOIN campaign_contacts cc ON cc.contact_id=p.id WHERE cc.campaign_id=1').fetchall() == [('0912345678',)]
    # Existing campaign add-data supports TXT and group snapshot together.
    post(client, '/campaigns/1/data', {'source_tag': 'VIP', 'consent': 'yes',
        'file': (io.BytesIO(b'0934567890'), 'data.txt')})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM campaign_contacts WHERE campaign_id=1').fetchone() == (3,)


def test_import_errors_rollback_and_tenant_isolation(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Private'})
    post(client, '/data-groups/import', {'group': 'Private', 'phones': '0901234567',
        'file': (io.BytesIO(b'phone,name\ninvalid,Bad'), 'bad.csv')})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)
    other = app.test_client()
    register(other, 'other')
    assert other.get('/data-groups?group=Private').status_code == 404
    assert post(other, '/data-groups/import', {'group': 'Private', 'phones': '0901234567'}).status_code == 404
    post(client, '/campaigns', {'name': 'Missing', 'message': 'Hi', 'data_mode': 'direct', 'phones': '0901234567'})
    post(client, '/campaigns', {'name': 'Empty', 'message': 'Hi', 'data_mode': 'direct'})
    post(client, '/campaigns', {'name': 'Missing group', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Secret'})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM campaigns').fetchone() == (0,)


def test_invalid_excel_and_large_import_are_atomic(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'VIP'})
    for content, filename in [(b'broken', 'bad.xlsx'), (b'x' * (2 * 1024 * 1024 + 1), 'big.csv'),
                              (b'phone\n' + b'0901234567\n' * 501, 'many.csv')]:
        response = post(client, '/data-groups/import', {'group': 'VIP', 'consent': 'yes',
            'phones': '0912345678', 'file': (io.BytesIO(content), filename)})
        assert response.status_code == 200
        with database(app) as db:
            assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)


def test_create_group_dialog_and_saved_preferences(app):
    client = app.test_client()
    register(client)
    page = client.get('/data-groups').text
    assert 'id="create-group-dialog"' in page
    assert 'Tạo nhóm data' in page
    response = post(client, '/data-groups', {'name': 'Tháng 10', 'data_type': 'phone', 'color': '#EF1234'})
    assert response.status_code == 200
    assert 'Số điện thoại' in response.text
    assert 'data-group-color="#ef1234"' in response.text
    with database(app) as db:
        assert db.execute('SELECT name,data_type,color FROM data_groups').fetchone() == ('Tháng 10', 'phone', '#ef1234')
    # The same group name must not silently reset its saved preferences.
    response = post(client, '/data-groups', {'name': 'Tháng 10', 'data_type': 'uid', 'color': '#000000'})
    assert 'Tên nhóm đã tồn tại' in response.text
    assert 'data-reopen="yes"' in response.text
    with database(app) as db:
        assert db.execute('SELECT data_type,color FROM data_groups').fetchall() == [('phone', '#ef1234')]


def test_invalid_group_preferences_and_csrf_are_rejected(app):
    client = app.test_client()
    register(client)
    for fields in [{'name': '  '}, {'name': 'Invalid', 'data_type': 'unknown'},
                   {'name': 'Invalid', 'color': 'red;display:none'}]:
        assert post(client, '/data-groups', fields).status_code == 200
    assert client.post('/data-groups', data={'name': 'Missing CSRF'}).status_code == 400
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM data_groups').fetchone() == (0,)


def test_legacy_groups_receive_metadata_without_losing_contacts(app):
    from app import create_app
    client = app.test_client()
    register(client)
    post(client, '/contacts', {'name': 'Lan', 'phone': '0901234567', 'tag': 'Legacy'})
    with database(app) as db:
        db.execute('DROP TABLE data_groups')
        db.execute('CREATE TABLE data_groups (id INTEGER PRIMARY KEY,user_id INTEGER NOT NULL,name TEXT NOT NULL,UNIQUE(user_id,name))')
        db.execute("INSERT INTO data_groups(user_id,name) VALUES (1,'Legacy')")
    create_app({'TESTING': True, 'DATABASE': app.config['DATABASE'], 'SECRET_KEY': 'test-only'})
    with database(app) as db:
        assert db.execute('SELECT name,data_type,color FROM data_groups').fetchone() == ('Legacy', 'all', '#6860f6')
        assert db.execute('SELECT name,tag FROM contacts').fetchone() == ('Lan', 'Legacy')


def test_group_detail_empty_state_filters_pagination_and_import_dates(app):
    client = app.test_client()
    register(client)
    response = post(client, '/data-groups', {'name': 'Details'})
    assert 'Không có data nào khớp bộ lọc.' in response.text
    assert 'Chưa có data để phân bố.' in response.text
    assert 'Đang dùng ở 0 chiến dịch' in response.text
    assert 'Trang 1/1' in response.text
    phones = '\n'.join(f'09{index:08}' for index in range(51))
    post(client, '/data-groups/import', {'group': 'Details', 'phones': phones, 'consent': 'yes'})
    response = client.get('/data-groups?group=Details')
    assert response.text.count('name="contact_ids"') == 50
    assert 'Trang 1/2' in response.text
    response = client.get('/data-groups?group=Details&p=2')
    assert response.text.count('name="contact_ids"') == 1
    assert 'Trang 2/2' in response.text
    assert 'Không có data nào khớp bộ lọc.' in client.get('/data-groups?group=Details&kind=uid').text
    assert 'Không có data nào khớp bộ lọc.' in client.get('/data-groups?group=Details&consent=no').text
    assert client.get('/data-groups?group=Details&kind=invalid').status_code == 400
    post(client, '/campaigns', {'name': 'Group use', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Details'})
    assert 'Đang dùng ở 1 chiến dịch' in client.get('/data-groups?group=Details').text
    with database(app) as db:
        created, latest, count = db.execute('SELECT created_at,last_import_at,imported_count FROM data_groups').fetchone()
        assert created and latest and count == 51


def test_export_group_selection_is_scoped_and_csv_safe(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Export'})
    post(client, '/data-groups/import', {'group': 'Export', 'consent': 'yes',
        'file': (io.BytesIO(b'phone,name\n0901234567,=1+1\n'), 'data.csv')})
    response = post(client, '/data-groups/export', {'group': 'Export', 'contact_ids': '1'})
    assert response.mimetype == 'text/csv'
    assert "'=1+1" in response.data.decode('utf-8-sig')
    assert '0901234567' in response.data.decode('utf-8-sig')
    assert post(client, '/data-groups/export', {'group': 'Wrong', 'contact_ids': '1'}).status_code == 404
    assert post(client, '/data-groups/export', {'group': 'Export'}).status_code == 400
    other = app.test_client()
    register(other, 'other')
    assert post(other, '/data-groups/export', {'group': 'Export', 'contact_ids': '1'}).status_code == 404


def test_account_dropdown_uses_owned_accounts_and_filters_campaign_data(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Accounts'})
    post(client, '/data-groups/import', {'group': 'Accounts', 'phones': '0901234567', 'consent': 'yes'})
    with database(app) as db:
        db.execute("INSERT INTO zalo_accounts(user_id,name) VALUES (1,'Kế Toán Viettax')")
        db.execute("INSERT INTO zalo_accounts(user_id,name) VALUES (1,'Admin kế toán')")
    response = client.get('/data-groups?group=Accounts')
    assert 'Kế Toán Viettax' in response.text and 'Admin kế toán' in response.text
    post(client, '/campaigns', {'name': 'Account campaign', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Accounts', 'account_id': '1'})
    response = client.get('/data-groups?group=Accounts&account_id=1')
    assert '<option value="1" selected>' in response.text
    assert 'name="contact_ids"' in response.text
    assert 'name="contact_ids"' not in client.get('/data-groups?group=Accounts&account_id=2').text
    assert 'name="contact_ids"' in client.get('/data-groups?group=Accounts').text
    other = app.test_client()
    register(other, 'other')
    assert 'Kế Toán Viettax' not in other.get('/data-groups').text
    assert other.get('/data-groups?account_id=1').status_code == 404
    assert client.get('/data-groups?account_id=invalid').status_code == 404


def test_import_preview_normalizes_without_saving_and_revalidates_on_save(app):
    from test_workflow import token
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Preview'})
    response = post(client, '/data-groups/import/preview', {'group': 'Preview', 'phones': '+84 901-234-567,0901234567\n0912345678'})
    assert response.json['total'] == 2 and response.json['added'] == 2
    assert [row['phone'] for row in response.json['rows']] == ['0901234567', '0912345678']
    assert response.json['rows'][0]['name'] == ''
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)
    post(client, '/data-groups/import', {'group': 'Preview', 'phones': '0901234567'})
    response = post(client, '/data-groups/import/preview', {'group': 'Preview', 'phones': '0901234567,0912345678'})
    assert response.json['existing'] == 1 and response.json['added'] == 1
    response = post(client, '/data-groups/import/preview', {'group': 'Preview', 'phones': 'not-a-phone'})
    assert response.status_code == 400 and response.json['error']
    post(client, '/data-groups/import', {'group': 'Preview', 'phones': 'invalid', 'consent': 'yes'})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (1,)
    assert client.post('/data-groups/import/preview', data={'group': 'Preview', 'phones': '0901234567'}).status_code == 400
    other = app.test_client()
    register(other, 'other')
    assert post(other, '/data-groups/import/preview', {'group': 'Preview', 'phones': '0901234567'}).status_code == 404


def test_excel_template_preview_keeps_personalization_and_skips_blank_formatted_rows(app):
    from openpyxl import load_workbook
    from pathlib import Path
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Excel Preview'})
    book = load_workbook(Path(app.root_path) / 'static/data-template.xlsx')
    book.active.append([])  # Template has formatting through row 501; write into its first data row.
    for column, value in enumerate(['0901234567', 'Lan', '', 'lan@example.com', 'Info A'], 1):
        book.active.cell(2, column, value)
    stream = io.BytesIO(); book.save(stream); stream.seek(0)
    response = post(client, '/data-groups/import/preview', {'group': 'Excel Preview', 'file': (stream, 'data.xlsx')})
    assert response.status_code == 200
    assert response.json['total'] == 1
    assert response.json['rows'][0]['info1'] == 'Info A'
    assert response.json['rows'][0]['email'] == 'lan@example.com'
    page = client.get('/data-groups?group=Excel+Preview').text
    assert 'Thêm data vào nhóm “Excel Preview”' in page
    assert 'Format chuẩn dữ liệu' in page
    assert 'File Excel/CSV của khách hàng' in page
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)


def test_customer_csv_column_mapping_inspection_preview_and_save(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Mapped'})
    raw = 'Ghi chú;Số điện thoại;Họ tên\nKhách VIP;+84901234567;Lan\n'.encode()
    fields = {'group': 'Mapped', 'import_mode': 'mapped', 'skip_header': 'yes',
              'column_name': '2', 'column_phone': '1', 'column_info3': '0'}
    response = post(client, '/data-groups/import/columns', {'group': 'Mapped', 'file': (io.BytesIO(raw), 'customer.csv')})
    assert response.json['columns'][1]['sample'] == 'Số điện thoại'
    response = post(client, '/data-groups/import/preview', {**fields, 'file': (io.BytesIO(raw), 'customer.csv')})
    assert response.status_code == 200
    assert response.json['rows'][0]['name'] == 'Lan'
    assert response.json['rows'][0]['phone'] == '0901234567'
    assert response.json['rows'][0]['info3'] == 'Khách VIP'
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)
    post(client, '/data-groups/import', {**fields, 'consent': 'yes', 'file': (io.BytesIO(raw), 'customer.csv')})
    with database(app) as db:
        assert db.execute('SELECT name,phone,info1,info3 FROM contacts').fetchone() == ('Lan', '0901234567', '', 'Khách VIP')


def test_customer_xlsx_mapping_without_header_keeps_first_contact(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'No header'})
    book = Workbook(); book.active.append(['Lan', 901234567, 'A']); book.active.append(['Bình', '0912345678', 'B'])
    stream = io.BytesIO(); book.save(stream)
    fields = {'group': 'No header', 'import_mode': 'mapped', 'column_name': '0', 'column_phone': '1', 'column_info1': '2'}
    response = post(client, '/data-groups/import/preview', {**fields, 'file': (io.BytesIO(stream.getvalue()), 'data.xlsx')})
    assert response.json['total'] == 2
    assert response.json['rows'][0]['phone'] == '0901234567'
    assert response.json['rows'][0]['info1'] == 'A'
    assert response.json['rows'][1]['name'] == 'Bình'


def test_customer_column_mapping_rejects_missing_duplicate_and_foreign_columns(app):
    client = app.test_client()
    register(client)
    post(client, '/data-groups', {'name': 'Validate'})
    raw = b'Name,Phone\nLan,0901234567\n'
    for mapping in [{}, {'column_phone': '99'}, {'column_phone': '-1'}, {'column_phone': '1', 'column_name': '1'}]:
        response = post(client, '/data-groups/import/preview', {'group': 'Validate', 'import_mode': 'mapped', 'skip_header': 'yes',
            **mapping, 'file': (io.BytesIO(raw), 'data.csv')})
        assert response.status_code == 400 and response.json['error']
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)
    other = app.test_client(); register(other, 'other')
    assert post(other, '/data-groups/import/columns', {'group': 'Validate', 'file': (io.BytesIO(raw), 'data.csv')}).status_code == 404
    assert client.post('/data-groups/import/columns', data={'group': 'Validate'}).status_code == 400


def test_template_legacy_xls_preview_and_save_with_vietnamese_headers(app):
    from pathlib import Path
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Legacy Excel'})
    data = (Path(__file__).parent / 'fixtures/template-legacy.xls').read_bytes()
    response = post(client, '/data-groups/import/preview', {'group': 'Legacy Excel', 'file': (io.BytesIO(data), 'template.xls')})
    assert response.status_code == 200
    assert response.json['rows'][0]['phone'] == '0901234567'
    assert response.json['rows'][0]['name'] == 'Lan'
    assert response.json['rows'][0]['info1'] == 'Thông tin A'
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)
    post(client, '/data-groups/import', {'group': 'Legacy Excel', 'file': (io.BytesIO(data), 'template.xls'), 'consent': 'yes'})
    with database(app) as db:
        assert db.execute('SELECT name,phone,info5 FROM contacts').fetchone() == ('Lan', '0901234567', 'E')
    response = post(client, '/data-groups/import/preview', {'group': 'Legacy Excel', 'import_mode': 'mapped',
        'skip_header': 'yes', 'column_phone': '1', 'column_name': '0', 'file': (io.BytesIO(data), 'customer.xls')})
    assert response.json['existing'] == 1


def test_template_rejects_invalid_xls_without_saving(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Invalid XLS'})
    response = post(client, '/data-groups/import/preview', {'group': 'Invalid XLS', 'file': (io.BytesIO(b'broken'), 'broken.xls')})
    assert response.status_code == 400 and response.json['error']
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)


def test_raw_excel_preview_preserves_columns_before_mapping_and_limits_rows(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Raw Preview'})
    book = Workbook(); sheet = book.active
    sheet.append(['MST', 'Ngày', 'Công ty', '', 'Thành phố', '', 'SĐT', 'Email'])
    for index in range(7):
        sheet.append([f'031234567{index}', '2026-10-19', f'Công ty {index}', '', 'Hồ Chí Minh', '', '0384.410.323', f'{index}@example.com'])
    stream = io.BytesIO(); book.save(stream); raw = stream.getvalue()
    response = post(client, '/data-groups/import/columns', {'group': 'Raw Preview', 'file': (io.BytesIO(raw), 'customer.xlsx')})
    assert response.status_code == 200
    assert response.json['row_count'] == 8
    assert len(response.json['columns']) == 8
    assert len(response.json['preview_rows']) == 6  # Header plus first five contacts.
    assert response.json['preview_rows'][1][3] == ''
    assert response.json['preview_rows'][1][6] == '0384.410.323'
    response = post(client, '/data-groups/import/preview', {'group': 'Raw Preview', 'import_mode': 'mapped',
        'skip_header': 'yes', 'column_name': '2', 'column_phone': '6', 'column_info1': '0', 'file': (io.BytesIO(raw), 'customer.xlsx')})
    assert response.status_code == 200
    assert response.json['rows'][0]['phone'] == '0384410323'
    assert response.json['rows'][0]['name'] == 'Công ty 6'  # Duplicate phone keeps the last row.
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)


def test_mapped_import_reports_bad_rows_without_blocking_valid_preview(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Mixed'})
    raw = b'Phone,Name\n0384.410.323,Lan\n,Binh\nPhone,Header\n099,Invalid\n0912345678,An\n'
    fields = {'group': 'Mixed', 'import_mode': 'mapped', 'skip_header': 'yes', 'column_phone': '0', 'column_name': '1'}
    response = post(client, '/data-groups/import/preview', {**fields, 'file': (io.BytesIO(raw), 'mixed.csv')})
    assert response.status_code == 200
    assert response.json['total'] == 2 and response.json['invalid'] == 3
    assert [row['row'] for row in response.json['invalid_rows']] == [3, 4, 5]
    assert response.json['invalid_rows'][2]['phone'] == '099'
    assert response.json['rows'][0]['phone'] == '0384410323'
    # Partial imports require explicit confirmation, even when posting directly.
    post(client, '/data-groups/import', {**fields, 'file': (io.BytesIO(raw), 'mixed.csv')})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)
    post(client, '/data-groups/import', {**fields, 'skip_invalid': 'yes', 'file': (io.BytesIO(raw), 'mixed.csv')})
    with database(app) as db:
        assert db.execute('SELECT phone FROM contacts ORDER BY id').fetchall() == [('0384410323',), ('0912345678',)]
        assert db.execute('SELECT imported_count FROM data_groups').fetchone() == (2,)


def test_all_invalid_mapped_rows_are_reported_and_cannot_be_saved(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'All invalid'})
    fields = {'group': 'All invalid', 'import_mode': 'mapped', 'column_phone': '0'}
    raw = b'not-a-phone\n12345\n'
    response = post(client, '/data-groups/import/preview', {**fields, 'file': (io.BytesIO(raw), 'invalid.csv')})
    assert response.status_code == 200 and response.json['total'] == 0 and response.json['invalid'] == 2
    post(client, '/data-groups/import', {**fields, 'skip_invalid': 'yes', 'file': (io.BytesIO(raw), 'invalid.csv')})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (0,)


def test_group_table_tracks_file_source_and_preserves_it_for_duplicates(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Sources'})
    data = b'phone\n0866381540\n'
    response = post(client, '/data-groups/import', {'group': 'Sources', 'file': (io.BytesIO(data), 'CSKH (3).csv')})
    assert 'Tải file lên · CSKH (3).csv' in response.text
    for title in ('Tên hiển thị','UID','Link','Tài khoản','Loại','Trạng thái','Nguồn data'):
        assert f'<th>{title}</th>' in response.text
    assert 'group-contact-name' in response.text
    post(client, '/data-groups/import', {'group': 'Sources', 'file': (io.BytesIO(data), 'Different.csv')})
    with database(app) as db:
        assert db.execute('SELECT data_source FROM contacts').fetchone() == ('Tải file lên · CSKH (3).csv',)


def test_edit_group_keeps_id_data_and_campaign_audience(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Before'})
    post(client, '/data-groups/import', {'group': 'Before', 'phones': '0901234567', 'consent': 'yes'})
    post(client, '/campaigns', {'name': 'Keep', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Before'})
    with database(app) as db:
        original_id = db.execute('SELECT id FROM data_groups WHERE name=?', ('Before',)).fetchone()[0]
    response = post(client, '/data-groups/edit', {'original_group': 'Before', 'name': 'After', 'data_type': 'phone', 'color': '#123456'})
    assert response.status_code == 200
    with database(app) as db:
        assert db.execute('SELECT id,name,color FROM data_groups WHERE id=?', (original_id,)).fetchone() == (original_id, 'After', '#123456')
        assert db.execute('SELECT phone,tag FROM contacts').fetchone() == ('0901234567', 'After')
        assert db.execute('SELECT tag FROM campaigns').fetchone() == ('After',)
    other = app.test_client(); register(other, 'other')
    assert post(other, '/data-groups/edit', {'original_group': 'After', 'name': 'Stolen'}).status_code == 404
    post(client, '/data-groups', {'name': 'Exists'})
    post(client, '/data-groups/edit', {'original_group': 'After', 'name': 'Exists'})
    with database(app) as db:
        assert db.execute('SELECT tag FROM contacts').fetchone() == ('After',)


def test_delete_group_requires_confirmation_and_preserves_contacts_campaigns(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Archive'})
    post(client, '/data-groups/import', {'group': 'Archive', 'phones': '0901234567', 'consent': 'yes'})
    post(client, '/campaigns', {'name': 'Keep', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Archive'})
    assert 'giữ nguyên' in client.get('/data-groups/delete?group=Archive').text
    assert post(client, '/data-groups/delete', {'group': 'Archive'}).status_code == 400
    post(client, '/data-groups/delete', {'group': 'Archive', 'confirm_delete': 'yes'})
    assert client.get('/data-groups?group=Archive').status_code == 404
    assert client.get('/campaigns/1').status_code == 200
    assert post(client, '/data-groups/import', {'group': 'Archive', 'phones': '0912345678'}).status_code == 404
    with database(app) as db:
        assert db.execute('SELECT phone FROM contacts').fetchall() == [('0901234567',)]
        assert db.execute('SELECT deleted_at FROM data_groups WHERE name=?', ('Archive',)).fetchone()[0]
        assert db.execute('SELECT count(*) FROM campaigns').fetchone() == (1,)
    other = app.test_client(); register(other, 'other')
    assert other.get('/data-groups/delete?group=Archive').status_code == 404


def test_duplicate_group_snapshots_data_without_duplicating_customer_profiles(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Original', 'color': '#123456', 'data_type': 'phone'})
    raw = b'phone,name,email,info1\n0901234567,Lan,lan@example.com,A\n0912345678,An,an@example.com,B\n'
    post(client, '/data-groups/import', {'group': 'Original', 'consent': 'yes', 'file': (io.BytesIO(raw), 'original.csv')})
    with database(app) as db:
        db.execute("UPDATE contacts SET consent=0 WHERE phone='0912345678'")
    assert '2 data' in client.get('/data-groups/duplicate?group=Original').text
    response = post(client, '/data-groups/duplicate', {'group': 'Original', 'name': 'Copy'})
    assert response.status_code == 200 and 'lan@example.com' not in response.text  # Email remains in the customer profile.
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (2,)
        assert db.execute("SELECT color,data_type FROM data_groups WHERE name='Copy'").fetchone() == ('#123456', 'phone')
        assert db.execute('SELECT count(DISTINCT id) FROM data_groups').fetchone() == (2,)
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='Copy'").fetchone() == (2,)
        assert db.execute("SELECT tag,email,info1 FROM contacts WHERE phone='0901234567'").fetchone() == ('Original', 'lan@example.com', 'A')
    response = post(client, '/data-groups/export', {'group': 'Copy', 'contact_ids': ['1','2']})
    assert response.mimetype == 'text/csv' and 'lan@example.com' in response.text
    # Appending to the original does not alter the copied snapshot.
    post(client, '/data-groups/import', {'group': 'Original', 'phones': '0923456789', 'consent': 'yes'})
    with database(app) as db:
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='Copy'").fetchone() == (2,)
    post(client, '/data-groups/duplicate', {'group': 'Copy', 'name': 'Copy again'})
    with database(app) as db:
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='Copy again'").fetchone() == (2,)
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (3,)


def test_copied_group_works_for_campaign_audience_snapshot_rename_and_archive(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Original'})
    post(client, '/data-groups/import', {'group': 'Original', 'phones': '0901234567', 'consent': 'yes'})
    post(client, '/data-groups/duplicate', {'group': 'Original', 'name': 'Copy'})
    post(client, '/campaigns', {'name': 'Copied audience', 'message': 'Hi', 'data_mode': 'group', 'tag': 'Copy'})
    post(client, '/campaigns/1')  # Existing simulation exercises the copied audience without real Zalo calls.
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM deliveries WHERE campaign_id=1').fetchone() == (1,)
    post(client, '/campaigns', {'name': 'Snapshot', 'message': 'Hi', 'data_mode': 'direct', 'source_tags': 'Copy', 'phones_consent': 'yes'})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM campaign_contacts WHERE campaign_id=2').fetchone() == (1,)
    post(client, '/campaigns', {'name': 'Append', 'message': 'Hi'})
    post(client, '/campaigns/3/data', {'source_tag': 'Copy', 'consent': 'yes'})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM campaign_contacts WHERE campaign_id=3').fetchone() == (1,)
    post(client, '/data-groups/edit', {'original_group': 'Copy', 'name': 'Renamed', 'color': '#123456'})
    with database(app) as db:
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='Renamed'").fetchone() == (1,)
        assert db.execute('SELECT tag FROM campaigns WHERE id=1').fetchone() == ('Renamed',)
        assert db.execute('SELECT tag FROM contacts').fetchone() == ('Original',)
    post(client, '/data-groups/delete', {'group': 'Original', 'confirm_delete': 'yes'})
    assert client.get('/data-groups?group=Renamed').status_code == 200
    assert client.get('/campaigns/1').status_code == 200
    with database(app) as db:
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='Renamed'").fetchone() == (1,)


def test_duplicate_group_validates_name_authentication_and_ownership(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name': 'Private'})
    for name in ['', 'Private', 'x'*81]:
        post(client, '/data-groups/duplicate', {'group': 'Private', 'name': name})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM data_groups').fetchone() == (1,)
    assert client.post('/data-groups/duplicate', data={'group':'Private','name':'No CSRF'}).status_code == 400
    other = app.test_client(); register(other, 'other')
    assert other.get('/data-groups/duplicate?group=Private').status_code == 404
    assert post(other, '/data-groups/duplicate', {'group':'Private','name':'Stolen'}).status_code == 404
    post(client, '/data-groups/duplicate', {'group':'Private','name':'Empty copy'})
    assert 'Không có data nào khớp bộ lọc.' in client.get('/data-groups?group=Empty+copy').text


def test_reimport_fills_missing_source_and_preserves_customer_information(app):
    client = app.test_client(); register(client)
    post(client, '/contacts', {'name': 'Existing Lan', 'phone': '0901234567', 'tag': 'Legacy', 'consent': ''})
    response = client.get('/data-groups?group=Legacy')
    assert 'Chưa ghi nhận nguồn' in response.text
    assert client.get('/data-groups/source-history/1?group=Legacy').json['sources'] == []
    raw = b'phone,name\n0901234567,Changed name\n'
    post(client, '/data-groups/import', {'group': 'Legacy', 'consent': 'yes', 'file': (io.BytesIO(raw), 'CSKH (3).csv')})
    with database(app) as db:
        name, consent, source = db.execute('SELECT name,consent,data_source FROM contacts').fetchone()
        assert (name, consent) == ('Existing Lan', 0)
        assert source == 'Tải file lên · CSKH (3).csv'
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (1,)
    response = client.get('/data-groups/source-history/1?group=Legacy')
    event = response.json['sources'][0]
    assert event['filename'] == 'CSKH (3).csv' and event['current']
    assert event['created_at'] and event['group_id'] and event['batch_id']
    post(client, '/data-groups/import', {'group': 'Legacy', 'file': (io.BytesIO(raw), 'Second.csv')})
    sources = client.get('/data-groups/source-history/1?group=Legacy').json['sources']
    assert len(sources) == 2
    assert sources[0]['filename'] == 'Second.csv' and not sources[0]['current']
    assert sources[1]['filename'] == 'CSKH (3).csv' and sources[1]['current']


def test_source_history_is_shared_by_group_copies_but_scoped_to_owner(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name':'Sources'})
    post(client, '/data-groups/import', {'group':'Sources', 'file':(io.BytesIO(b'phone\n0901234567\n'),'source.csv')})
    post(client, '/data-groups/duplicate', {'group':'Sources','name':'Copy'})
    assert client.get('/data-groups/source-history/1?group=Copy').json['sources'][0]['filename'] == 'source.csv'
    post(client, '/data-groups', {'name':'Other group'})
    assert client.get('/data-groups/source-history/1?group=Other+group').status_code == 404
    other = app.test_client(); register(other, 'other')
    post(other, '/data-groups', {'name':'Sources'})
    assert other.get('/data-groups/source-history/1?group=Sources').status_code == 404


def test_old_source_labels_migrate_without_inventing_date_or_batch(app):
    from app import create_app
    client = app.test_client(); register(client)
    post(client, '/contacts', {'name':'Lan','phone':'0901234567','tag':'Old'})
    with database(app) as db:
        db.execute("UPDATE contacts SET data_source='Tải file lên · old.xlsx'")
    for _ in range(2):
        create_app({'TESTING':True,'DATABASE':app.config['DATABASE'],'SECRET_KEY':'test-only'})
    source = client.get('/data-groups/source-history/1?group=Old').json['sources'][0]
    assert source['filename'] == 'old.xlsx' and source['current']
    assert source['created_at'] is None and source['batch_id'] is None
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contact_data_sources').fetchone() == (1,)


def test_import_131_existing_profiles_into_new_group_with_invalid_rows(app):
    client = app.test_client(); register(client)
    post(client, '/data-groups', {'name':'Original'})
    post(client, '/data-groups', {'name':'T100'})
    phones = [f'09{index:08}' for index in range(131)]
    with database(app) as db:
        db.executemany('INSERT INTO contacts(user_id,name,phone,tag,consent,email,info1) VALUES (1,?,?,?,0,?,?)',
                       [(f'Old name {index}',phone,'Original',f'old{index}@example.com','Keep') for index,phone in enumerate(phones)])
    stream = io.StringIO()
    import csv
    writer = csv.writer(stream); writer.writerow(['phone','name'])
    for phone in phones: writer.writerow([phone,'Uploaded name'])
    for bad in ['02822106868','099','bad-phone','13910216623','0933655998 - 0903066147']: writer.writerow([bad,'Bad'])
    raw = stream.getvalue().encode()
    fields = {'group':'T100','import_mode':'mapped','skip_header':'yes','column_phone':'0','column_name':'1'}
    response = post(client, '/data-groups/import/preview', {**fields,'file':(io.BytesIO(raw),'CSKH (3).csv')})
    assert response.json['total'] == 131 and response.json['added'] == 131
    assert response.json['existing'] == 0 and response.json['invalid'] == 5
    # Confirmation is still required for the five rejected rows.
    post(client, '/data-groups/import', {**fields,'file':(io.BytesIO(raw),'CSKH (3).csv')})
    with database(app) as db:
        assert db.execute("SELECT count(*) FROM contact_group_memberships WHERE group_name='T100'").fetchone() == (0,)
    result = post(client, '/data-groups/import', {**fields,'skip_invalid':'yes','consent':'yes','file':(io.BytesIO(raw),'CSKH (3).csv')})
    assert 'Đã thêm 131 số vào nhóm' in result.text
    assert '131 data trong nhóm' in result.text
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (131,)
        for group in ['Original','T100']:
            assert db.execute('SELECT count(*) FROM contact_group_memberships WHERE group_name=?',(group,)).fetchone() == (131,)
        assert db.execute('SELECT name,tag,consent,email,info1 FROM contacts ORDER BY id LIMIT 1').fetchone() == ('Old name 0','Original',0,'old0@example.com','Keep')
        assert db.execute("SELECT count(*) FROM contacts WHERE data_source='Tải file lên · CSKH (3).csv'").fetchone() == (131,)
    assert client.get('/data-groups/source-history/1?group=T100').json['sources'][0]['filename'] == 'CSKH (3).csv'
    response = post(client, '/data-groups/import/preview', {**fields,'file':(io.BytesIO(raw),'CSKH (3).csv')})
    assert response.json['added'] == 0 and response.json['existing'] == 131
    post(client, '/data-groups/import', {**fields,'skip_invalid':'yes','file':(io.BytesIO(raw),'CSKH (3).csv')})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM data_group_contacts').fetchone() == (131,)
        assert db.execute("SELECT imported_count FROM data_groups WHERE name='T100'").fetchone() == (131,)


def test_linking_existing_contacts_never_crosses_account_ownership(app):
    first = app.test_client(); register(first)
    post(first, '/data-groups', {'name':'Private'})
    post(first, '/data-groups/import', {'group':'Private','phones':'0901234567','consent':'yes'})
    second = app.test_client(); register(second,'other')
    post(second, '/data-groups', {'name':'Target'})
    post(second, '/data-groups/import', {'group':'Target','phones':'0901234567'})
    with database(app) as db:
        assert db.execute('SELECT count(*) FROM contacts').fetchone() == (2,)
        assert db.execute("SELECT user_id FROM contact_group_memberships WHERE group_name='Target'").fetchall() == [(2,)]
        assert db.execute("SELECT consent FROM contacts WHERE user_id=1").fetchone() == (1,)
