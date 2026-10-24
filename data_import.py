"""Bounded, validated contact imports shared by campaigns and data groups."""
import csv
import io
import unicodedata
from pathlib import Path
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError


def read_contacts(upload, parse_phones):
    if not upload or not upload.filename:
        return {}
    data = upload.read(2 * 1024 * 1024 + 1)
    if len(data) > 2 * 1024 * 1024:
        raise ValueError('Tệp data không vượt quá 2 MB.')
    suffix = Path(upload.filename).suffix.lower()
    try:
        if suffix == '.txt':
            return {phone: {'phone': phone} for phone in parse_phones(data.decode('utf-8-sig'))}
        if suffix == '.csv':
            reader = csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
            if reader.fieldnames:
                reader.fieldnames = [canonical_header(cell) for cell in reader.fieldnames]
            if not reader.fieldnames or 'phone' not in reader.fieldnames:
                raise ValueError('Tệp cần cột phone; tùy chọn name, zalo_id, email, info1…info5.')
            rows = reader
        elif suffix in ('.xlsx', '.xls'):
            from werkzeug.datastructures import FileStorage
            table = read_sheet_rows(FileStorage(stream=io.BytesIO(data), filename=upload.filename))
            headers = [canonical_header(cell) for cell in table[0]]
            if 'phone' not in headers:
                raise ValueError('Excel cần hàng đầu có cột phone hoặc Số điện thoại.')
            rows = (dict(zip(headers, row)) for row in table[1:])
        else:
            raise ValueError('Chọn tệp TXT, CSV UTF-8 hoặc Excel .xlsx/.xls.')
        result = {}
        for index, row in enumerate(rows):
            if index >= 500:
                raise ValueError('Tệp tối đa 500 dòng.')
            raw = str(row.get('phone') or '').strip()
            # Spreadsheet numeric cells drop the initial zero of Vietnamese numbers.
            if raw.isdigit() and len(raw) == 9:
                raw = '0' + raw
            phones = parse_phones(raw)
            if len(phones) != 1:
                raise ValueError(f'Dòng {index + 2}: cần đúng một số điện thoại.')
            phone = phones[0]
            clean = {key: str(row.get(key) or '').strip() for key in ('name', 'zalo_id', 'email', 'info1', 'info2', 'info3', 'info4', 'info5')}
            if len(clean['name']) > 100 or len(clean['zalo_id']) > 100 or len(clean['email']) > 254 or any(len(clean[key]) > 500 for key in ('info1','info2','info3','info4','info5')):
                raise ValueError(f'Dòng {index + 2}: thông tin vượt độ dài cho phép.')
            result[phone] = dict(clean, phone=phone)
        return result
    except (UnicodeError, csv.Error, BadZipFile, StopIteration, KeyError, TypeError, ParseError, OSError) as error:
        raise ValueError('Tệp không hợp lệ. Dùng TXT/CSV UTF-8 hoặc Excel .xlsx có cột phone.') from error


def default_contact_consent(form, checked_value):
    """Default new records to enabled while honoring an explicitly unchecked form."""
    default = '' if form.get('consent_present') == 'yes' else checked_value
    return int(form.get('consent', default) == checked_value)


def insert_contact(connection, user_id, phone, row, tag='Khách hàng', consent=1):
    connection.execute('''INSERT INTO contacts(user_id,name,phone,tag,consent,zalo_id,email,info1,info2,info3,info4,info5)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(user_id,phone) DO NOTHING''',
        (user_id, row.get('name') or 'Khách hàng ' + phone, phone, tag, consent,
         *[row.get(key, '') for key in ('zalo_id','email','info1','info2','info3','info4','info5')]))


def read_sheet_rows(upload):
    """Read bounded customer sheets without assuming column names."""
    if not upload or not upload.filename:
        raise ValueError('Chọn file Excel/CSV của khách hàng.')
    data = upload.read(2 * 1024 * 1024 + 1)
    if len(data) > 2 * 1024 * 1024:
        raise ValueError('Tệp data không vượt quá 2 MB.')
    suffix = Path(upload.filename).suffix.lower()
    book = None
    try:
        if suffix == '.csv':
            text = data.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            source = csv.reader(io.StringIO(text), dialect)
        elif suffix == '.xls':
            import xlrd
            try:
                legacy = xlrd.open_workbook(file_contents=data, on_demand=True)
                try:
                    sheet = legacy.sheet_by_index(0)
                    if sheet.nrows > 5000 or sheet.ncols > 100:
                        raise ValueError('Tệp tối đa 5000 dòng kể cả dòng trống và 100 cột.')
                    source = [sheet.row_values(index) for index in range(sheet.nrows)]
                finally:
                    legacy.release_resources()
            except xlrd.XLRDError as error:
                raise ValueError('Tệp Excel .xls không hợp lệ.') from error
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            from zipfile import ZipFile
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Tệp Excel giải nén quá lớn.')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            source = book.active.iter_rows(values_only=True)
        else:
            raise ValueError('Chọn CSV UTF-8 hoặc Excel .xlsx/.xls.')
        rows = []
        for index, row in enumerate(source):
            if index >= 5000:
                raise ValueError('Tệp quá nhiều dòng trống; hãy xóa các dòng thừa.')
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            if len(row) > 100:
                raise ValueError('Tệp tối đa 100 cột.')
            rows.append(['' if cell is None else str(int(cell)) if isinstance(cell, float) and cell.is_integer() else str(cell).strip() for cell in row])
            if len(rows) > 501:
                raise ValueError('Tệp tối đa 500 dòng data và một dòng tiêu đề.')
        if not rows:
            raise ValueError('Tệp chưa có dữ liệu.')
        return rows
    except (UnicodeError, csv.Error, BadZipFile, KeyError, TypeError, ParseError, OSError) as error:
        raise ValueError('Tệp Excel/CSV không hợp lệ hoặc không phải UTF-8.') from error
    finally:
        if book:
            book.close()


def read_mapped_contacts(upload, parse_phones, mapping, skip_header, report=False):
    from werkzeug.datastructures import FileStorage
    rows = read_sheet_rows(upload)
    width = max(len(row) for row in rows)
    fields = ('name','phone','info1','info2','info3','info4','info5')
    chosen = {}
    for field in fields:
        value = mapping.get(field, '')
        if value == '':
            continue
        if not value.isdigit() or int(value) >= width:
            raise ValueError('Cột được chọn không có trong tệp.')
        chosen[field] = int(value)
    if 'phone' not in chosen:
        raise ValueError('Cần chọn cột Số điện thoại.')
    if len(set(chosen.values())) != len(chosen):
        raise ValueError('Mỗi trường cần chọn một cột khác nhau.')
    data_rows = rows[1:] if skip_header else rows
    if not data_rows or len(data_rows) > 500:
        raise ValueError('Tệp cần từ 1 đến 500 dòng data sau khi bỏ tiêu đề.')
    valid = {}
    invalid = []
    for index, row in enumerate(data_rows):
        values = [row[chosen[field]] if field in chosen and chosen[field] < len(row) else '' for field in fields]
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(fields)
        writer.writerow(values)
        try:
            valid.update(read_contacts(FileStorage(stream=io.BytesIO(stream.getvalue().encode('utf-8')), filename='mapped.csv'), parse_phones))
        except ValueError as error:
            invalid.append({'row': index + (2 if skip_header else 1), 'phone': values[1][:100], 'reason': str(error).removeprefix('Dòng 2: ')})
    if report:
        return valid, invalid
    if invalid:
        first = invalid[0]
        raise ValueError(f"Dòng {first['row']}, SĐT {first['phone']!r}: {first['reason']}")
    return valid


def canonical_header(value):
    normalized = ''.join(char for char in unicodedata.normalize('NFD', str(value or '').strip().lower()) if unicodedata.category(char) != 'Mn').replace('đ', 'd')
    aliases = {'ten': 'name', 'ho ten': 'name', 'sdt': 'phone', 'so dien thoai': 'phone', 'phone number': 'phone', 'uid': 'zalo_id'}
    if normalized.startswith('info '):
        normalized = normalized.replace(' ', '')
    return aliases.get(normalized, normalized)
